"""Tests for PowerPoint slide text, table, and image extraction."""

from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.main import create_app
from app.services.chunking import TextChunker
from app.services.powerpoint_processing import (
    POWERPOINT_MIME_TYPE,
    InvalidPowerPointTypeError,
    MalformedPowerPointError,
    PowerPointDocumentProcessor,
    PowerPointLimitError,
    PowerPointTooLargeError,
)


def save_presentation(path: Path, presentation: Presentation) -> Path:
    """Save a presentation fixture and return its path."""
    presentation.save(path)
    return path


def test_text_slides_use_one_based_slide_numbers_and_feed_chunker(tmp_path: Path) -> None:
    """Text and titles are extracted in slide order with one-based numbering."""
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Cell Biology"
    first.placeholders[1].text = "Cells are the basic unit of life."
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = second.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "Organelles have specialized functions."
    path = save_presentation(tmp_path / "biology.pptx", presentation)

    result = PowerPointDocumentProcessor().process_path(
        path,
        "biology.pptx",
        document_id="doc-bio",
        source_id="upload-19",
    )

    assert [slide.slide_number for slide in result.slides] == [1, 2]
    assert result.slides[0].title == "Cell Biology"
    assert "Cells are the basic unit of life." in result.slides[0].text
    assert result.slides[1].title is None
    assert result.slides[1].text == "Organelles have specialized functions."
    assert result.source_filename == "biology.pptx"
    assert result.source_type == POWERPOINT_MIME_TYPE
    chunk = TextChunker().chunk(result.segments)[0]
    assert chunk.document_id == "doc-bio"
    assert chunk.source_id == "upload-19"
    assert chunk.slide_number == 1
    assert chunk.heading == "Cell Biology"


def test_table_text_is_extracted_in_row_order(tmp_path: Path) -> None:
    """Table cells are preserved as tab-separated rows on their source slide."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Example"
    table.cell(1, 0).text = "Mitosis"
    table.cell(1, 1).text = "Cell division"
    path = save_presentation(tmp_path / "table.pptx", presentation)

    result = PowerPointDocumentProcessor().process_path(path, "table.pptx")

    assert result.slides[0].slide_number == 1
    assert result.slides[0].text == "Topic\tExample\nMitosis\tCell division"


def test_image_assets_keep_their_slide_relationship_without_ocr(tmp_path: Path) -> None:
    """Images are extracted as bytes with slide metadata, not interpreted as text."""
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 24, 24), 0)
    pixmap.clear_with(160)
    image_bytes = pixmap.tobytes("png")

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    image_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    image_slide.shapes.add_picture(BytesIO(image_bytes), Inches(1), Inches(1))
    path = save_presentation(tmp_path / "image.pptx", presentation)

    result = PowerPointDocumentProcessor().process_path(path, "image.pptx")

    assert result.slides[1].slide_number == 2
    assert result.slides[1].status == "no_text"
    assert result.slides[1].image_count == 1
    assert len(result.images) == 1
    assert result.images[0].slide_number == 2
    assert result.images[0].content_type == "image/png"
    assert result.images[0].data == image_bytes
    assert result.images[0].asset_id
    assert result.status == "empty"


def test_empty_slide_without_title_is_reported_without_failure(tmp_path: Path) -> None:
    """Slides with neither title nor text are retained with an explicit status."""
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = save_presentation(tmp_path / "blank-slide.pptx", presentation)

    result = PowerPointDocumentProcessor().process_path(path, "blank-slide.pptx")

    assert result.status == "empty"
    assert result.slides[0].title is None
    assert result.slides[0].text == ""
    assert result.slides[0].status == "no_text"
    assert result.slides[0].slide_number == 1


def test_invalid_or_malformed_pptx_is_rejected(tmp_path: Path) -> None:
    """Wrong extensions and corrupt PPTX packages return typed safe errors."""
    text_path = tmp_path / "notes.txt"
    text_path.write_text("not a deck", encoding="utf-8")
    malformed_path = tmp_path / "broken.pptx"
    malformed_path.write_bytes(b"PK\x03\x04not a valid OOXML package")

    processor = PowerPointDocumentProcessor()
    with pytest.raises(InvalidPowerPointTypeError):
        processor.process_path(text_path, "notes.txt")
    with pytest.raises(MalformedPowerPointError):
        processor.process_path(malformed_path, "broken.pptx")


def test_upload_and_processing_limits_are_enforced(tmp_path: Path) -> None:
    """Upload size, expanded package size, and slide count have configured limits."""
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = save_presentation(tmp_path / "one-slide.pptx", presentation)

    with pytest.raises(PowerPointTooLargeError):
        PowerPointDocumentProcessor(max_upload_size_bytes=1).process_path(path, "one-slide.pptx")
    with pytest.raises(PowerPointTooLargeError):
        PowerPointDocumentProcessor(max_processing_size_bytes=1).process_path(path, "one-slide.pptx")

    assert PowerPointDocumentProcessor(max_slides=1).process_path(path, "one-slide.pptx").slide_count == 1
    two_slide_deck = Presentation()
    two_slide_deck.slides.add_slide(two_slide_deck.slide_layouts[6])
    two_slide_deck.slides.add_slide(two_slide_deck.slide_layouts[6])
    two_slide_path = save_presentation(tmp_path / "two-slides.pptx", two_slide_deck)
    with pytest.raises(PowerPointLimitError):
        PowerPointDocumentProcessor(max_slides=1).process_path(two_slide_path, "two-slides.pptx")


def test_pptx_upload_route_returns_slide_metadata() -> None:
    """The document API accepts PPTX while leaving PDF response shape separate."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "API response"
    slide.placeholders[1].text = "Parsed from the upload route."
    image_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), 0)
    pixmap.clear_with(180)
    image_slide.shapes.add_picture(BytesIO(pixmap.tobytes("png")), Inches(1), Inches(1))
    upload = BytesIO()
    presentation.save(upload)

    response = TestClient(create_app()).post(
        "/api/documents/upload/pptx",
        files={"file": ("api-test.pptx", upload.getvalue(), POWERPOINT_MIME_TYPE)},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["source_filename"] == "api-test.pptx"
    assert data["slides"][0]["slide_number"] == 1
    assert data["slides"][0]["title"] == "API response"
    assert data["slides"][0]["text"].startswith("API response")
    assert data["images"][0]["slide_number"] == 2
    assert "data" not in data["images"][0]
