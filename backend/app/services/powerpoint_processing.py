"""Page-aware PowerPoint text and embedded image extraction.

Basic parsing reads OOXML text and extracts image bytes only. It does not run
OCR or a vision model, and it does not interpret extracted image contents.
"""

from dataclasses import dataclass, field
import hashlib
import logging
from pathlib import Path
from typing import Any, Iterable, Literal, Protocol
from zipfile import BadZipFile, ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.chunking import SourceSegment
from app.services.document_processing import DocumentProcessingError

logger = logging.getLogger(__name__)
POWERPOINT_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class PowerPointProcessingError(DocumentProcessingError):
    """Base class for PowerPoint validation and extraction failures."""

    code = "powerpoint_processing_error"


class InvalidPowerPointTypeError(PowerPointProcessingError):
    """The supplied file is not a PPTX document."""

    code = "invalid_powerpoint_type"
    status_code = 415


class PowerPointTooLargeError(PowerPointProcessingError):
    """The PPTX exceeds the configured upload or expanded processing limit."""

    code = "powerpoint_too_large"
    status_code = 413


class PowerPointLimitError(PowerPointProcessingError):
    """The PPTX exceeds a configured slide, package, or text limit."""

    code = "powerpoint_processing_limit"
    status_code = 413


class MalformedPowerPointError(PowerPointProcessingError):
    """The PPTX package is malformed or unsupported."""

    code = "malformed_powerpoint"
    status_code = 422


@dataclass(frozen=True)
class ExtractedSlide:
    """Extracted slide content using a one-based source slide number."""

    source_filename: str
    source_type: Literal["application/vnd.openxmlformats-officedocument.presentationml.presentation"]
    slide_number: int
    title: str | None
    text: str
    status: Literal["extracted", "no_text"]
    image_count: int


@dataclass(frozen=True)
class ExtractedImage:
    """Raw image asset extracted from a slide, linked back to its source slide."""

    asset_id: str
    source_filename: str
    source_type: Literal["application/vnd.openxmlformats-officedocument.presentationml.presentation"]
    slide_number: int
    image_index: int
    asset_filename: str
    content_type: str
    size_bytes: int
    sha256: str
    data: bytes = field(repr=False)


@dataclass(frozen=True)
class PowerPointExtractionResult:
    """Structured PowerPoint extraction result and chunk-ready source segments."""

    source_filename: str
    source_type: Literal["application/vnd.openxmlformats-officedocument.presentationml.presentation"]
    status: Literal["completed", "completed_with_warnings", "empty"]
    slides: list[ExtractedSlide]
    images: list[ExtractedImage]
    slide_count: int
    extracted_char_count: int
    document_id: str | None = None
    source_id: str | None = None

    @property
    def segments(self) -> list[SourceSegment]:
        """Return slide text as generic segments accepted by the chunker."""
        return [
            SourceSegment(
                text=slide.text,
                source_filename=slide.source_filename,
                source_type=slide.source_type,
                document_id=self.document_id,
                source_id=self.source_id,
                slide_number=slide.slide_number,
                segment_id=f"slide-{slide.slide_number}",
                heading=slide.title,
                extraction_warnings=(
                    ["Slide contains image assets that were not analyzed by OCR or vision."]
                    if slide.image_count
                    else []
                ),
                is_incomplete=bool(slide.image_count),
            )
            for slide in self.slides
        ]


class SlideImageAnalyzer(Protocol):
    """Optional OCR or vision adapter; extraction never invokes this protocol."""

    async def analyze_image(self, image: ExtractedImage) -> str | None:
        """Return optional image-derived text or description for a slide asset."""
        ...


class PowerPointDocumentProcessor:
    """Extract slide text and image assets from PPTX files with safety limits."""

    def __init__(
        self,
        max_upload_size_bytes: int = 25 * 1024 * 1024,
        max_slides: int = 500,
        max_extracted_chars: int = 2_000_000,
        max_processing_size_bytes: int = 100 * 1024 * 1024,
        max_archive_entries: int = 10_000,
    ) -> None:
        """Configure upload, OOXML expansion, slide, and extracted-text limits."""
        if min(
            max_upload_size_bytes,
            max_slides,
            max_extracted_chars,
            max_processing_size_bytes,
            max_archive_entries,
        ) < 1:
            raise ValueError("PowerPoint processing limits must be positive")
        self.max_upload_size_bytes = max_upload_size_bytes
        self.max_slides = max_slides
        self.max_extracted_chars = max_extracted_chars
        self.max_processing_size_bytes = max_processing_size_bytes
        self.max_archive_entries = max_archive_entries

    def process_path(
        self,
        path: Path,
        filename: str,
        *,
        document_id: str | None = None,
        source_id: str | None = None,
    ) -> PowerPointExtractionResult:
        """Extract each slide from a temporary or trusted local PPTX path.

        Slide numbers are one-based in presentation order (first slide is 1),
        independent of any visible numbering text in the deck.
        """
        safe_filename = Path(filename.replace("\\", "/")).name
        if Path(safe_filename).suffix.lower() != ".pptx":
            raise InvalidPowerPointTypeError("Only .pptx files are accepted.")

        try:
            file_size = path.stat().st_size
        except OSError as exc:
            raise MalformedPowerPointError("The uploaded presentation could not be read.") from exc
        if file_size > self.max_upload_size_bytes:
            raise PowerPointTooLargeError("The PPTX exceeds the configured upload size limit.")
        if file_size == 0:
            raise MalformedPowerPointError("The uploaded PPTX is empty or invalid.")

        try:
            with path.open("rb") as file_handle:
                signature = file_handle.read(4)
            if signature != b"PK\x03\x04":
                raise InvalidPowerPointTypeError("The uploaded file is not a PPTX package.")

            self._validate_package(path)
            presentation = Presentation(str(path))
            slide_count = len(presentation.slides)
            if slide_count > self.max_slides:
                raise PowerPointLimitError("The presentation exceeds the configured slide limit.")

            slides: list[ExtractedSlide] = []
            images: list[ExtractedImage] = []
            extracted_char_count = 0
            for slide_index, slide in enumerate(presentation.slides, start=1):
                title_shape = slide.shapes.title
                title = self._clean_text(title_shape.text) if title_shape and title_shape.has_text_frame else None
                title = title or None
                text_parts: list[str] = []
                slide_images: list[ExtractedImage] = []

                for shape_index, shape in enumerate(self._walk_shapes(slide.shapes), start=1):
                    if shape.has_text_frame:
                        shape_text = self._clean_text(shape.text)
                        if shape_text:
                            text_parts.append(shape_text)
                    if shape.has_table:
                        table_text = self._extract_table(shape.table)
                        if table_text:
                            text_parts.append(table_text)
                    if shape.has_chart:
                        chart_text = self._extract_chart_text(shape.chart)
                        if chart_text:
                            text_parts.append(chart_text)
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        slide_images.append(
                            self._extract_image(
                                shape,
                                source_filename=safe_filename,
                                slide_number=slide_index,
                                image_index=len(slide_images) + 1,
                                shape_index=shape_index,
                            )
                        )

                slide_text = "\n".join(text_parts).strip()
                extracted_char_count += len(slide_text)
                if extracted_char_count > self.max_extracted_chars:
                    raise PowerPointLimitError("The PPTX exceeds the configured extracted text limit.")

                images.extend(slide_images)
                slides.append(
                    ExtractedSlide(
                        source_filename=safe_filename,
                        source_type=POWERPOINT_MIME_TYPE,
                        slide_number=slide_index,
                        title=title,
                        text=slide_text,
                        status="extracted" if slide_text else "no_text",
                        image_count=len(slide_images),
                    )
                )
        except DocumentProcessingError:
            raise
        except (BadZipFile, KeyError, OSError, ValueError, TypeError, AttributeError) as exc:
            logger.exception("PPTX extraction failed filename=%s", safe_filename)
            raise MalformedPowerPointError("The PPTX is malformed or could not be processed.") from exc
        except Exception as exc:
            logger.exception("PPTX extraction failed filename=%s", safe_filename)
            raise MalformedPowerPointError("The PPTX is malformed or could not be processed.") from exc

        if extracted_char_count == 0:
            status: Literal["completed", "completed_with_warnings", "empty"] = "empty"
        elif any(slide.status == "no_text" for slide in slides):
            status = "completed_with_warnings"
        else:
            status = "completed"
        return PowerPointExtractionResult(
            source_filename=safe_filename,
            source_type=POWERPOINT_MIME_TYPE,
            status=status,
            slides=slides,
            images=images,
            slide_count=len(slides),
            extracted_char_count=extracted_char_count,
            document_id=document_id,
            source_id=source_id,
        )

    def _validate_package(self, path: Path) -> None:
        """Reject malformed, oversized, or non-presentation OOXML ZIP packages."""
        try:
            with ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > self.max_archive_entries:
                    raise PowerPointLimitError("The PPTX contains too many package entries.")
                uncompressed_size = sum(entry.file_size for entry in entries)
                if uncompressed_size > self.max_processing_size_bytes:
                    raise PowerPointTooLargeError("The expanded PPTX exceeds the processing size limit.")
                names = {entry.filename for entry in entries}
                if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                    raise InvalidPowerPointTypeError("The ZIP package is not an OOXML presentation.")
        except BadZipFile as exc:
            raise MalformedPowerPointError("The PPTX package is malformed.") from exc

    @classmethod
    def _walk_shapes(cls, shapes: Iterable[Any]) -> list[Any]:
        """Yield shapes in slide order, including shapes nested in groups."""
        found: list[Any] = []
        for shape in shapes:
            found.append(shape)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                found.extend(cls._walk_shapes(shape.shapes))
        return found

    @staticmethod
    def _clean_text(value: str) -> str:
        """Normalize line endings and discard whitespace-only text."""
        return "\n".join(line.strip() for line in value.replace("\r", "\n").split("\n")).strip()

    @classmethod
    def _extract_table(cls, table: Any) -> str:
        """Return table cells in row order as tab-separated text."""
        rows = []
        for row in table.rows:
            cells = [cls._clean_text(cell.text) for cell in row.cells]
            rows.append("\t".join(cells).rstrip())
        return "\n".join(row for row in rows if row).strip()

    @classmethod
    def _extract_chart_text(cls, chart: Any) -> str:
        """Extract chart title, series labels, and category labels when present."""
        parts: list[str] = []
        if chart.has_title:
            title = cls._clean_text(chart.chart_title.text_frame.text)
            if title:
                parts.append(title)
        for series in chart.series:
            name = str(series.name).strip()
            if name:
                parts.append(name)
        for plot in chart.plots:
            try:
                for category in plot.categories:
                    label = str(category.label).strip()
                    if label:
                        parts.append(label)
            except (AttributeError, ValueError):
                # Some chart types have no category axis or expose nested data.
                continue
        return "\n".join(parts)

    @staticmethod
    def _extract_image(
        shape: Any,
        *,
        source_filename: str,
        slide_number: int,
        image_index: int,
        shape_index: int,
    ) -> ExtractedImage:
        """Extract raw image bytes without interpreting image contents."""
        image = shape.image
        data = image.blob
        digest = hashlib.sha256(data).hexdigest()
        asset_id = f"slide-{slide_number}-image-{image_index}-{digest[:12]}"
        return ExtractedImage(
            asset_id=asset_id,
            source_filename=source_filename,
            source_type=POWERPOINT_MIME_TYPE,
            slide_number=slide_number,
            image_index=image_index,
            asset_filename=image.filename or f"slide-{slide_number}-shape-{shape_index}.{image.ext}",
            content_type=image.content_type,
            size_bytes=len(data),
            sha256=digest,
            data=data,
        )
